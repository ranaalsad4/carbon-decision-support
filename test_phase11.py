"""Phase 11 - Supplier Decision & Engagement tests.   Run:  python -m pytest -v
Runs alongside test_workstream_b.py, test_workstream_c.py and phase4/test_phase4.py (full regression)."""
import copy
import hashlib
import json
import shutil
from math import isclose
from pathlib import Path

import openpyxl
import pytest

import calc_engine as ce
import run_workstream_c as rwc
import scenario_engine as se
import supplier_audit as sa
import supplier_carbon_engine as sce
import supplier_engagement_engine as see
import supplier_financial_engine as sfe
import supplier_recommendation_engine as sre
import supplier_scoring_engine as ssc
import supplier_screening_engine as sse
from supplier_data_loader import REGISTER, apply_offer_edits, inventory_by_supplier, load_register, offers_table
from supplier_models import UnitError, WeightError, normalise_ef, parse_number

HERE = Path(__file__).resolve().parent


@pytest.fixture(scope="module")
def data():
    return load_register()


@pytest.fixture(scope="module")
def res(data):
    return sa.analyse(copy.deepcopy(data))


@pytest.fixture
def d(data):
    """Fresh mutable copy for tests that change inputs."""
    return copy.deepcopy(data)


def _run(d, wtp="settings", weights=None, hypothetical=None):
    return sa.analyse(d, weights=weights, wtp=wtp, hypothetical=hypothetical)


def _sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def _mutate(tmp_path, sheet, row_key, col, value):
    """Copy of the register with one cell changed (header row 4, key in column A)."""
    p = tmp_path / "reg.xlsx"
    shutil.copy(REGISTER, p)
    wb = openpyxl.load_workbook(p)
    ws = wb[sheet]
    hdr = {c.value: c.column for c in ws[4] if c.value}
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == row_key:
            ws.cell(r, hdr[col]).value = value
            break
    else:
        raise KeyError(row_key)
    wb.save(p)
    return p


# ================================================================== 1 preservation (baseline untouched)
def test_inventory_scope_totals_preserved():
    audit, _ = ce.run()
    src = ce.by_source(audit)
    sc = {s["scope"]: s["emissions_t"] for s in ce.by_scope(src)}
    assert isclose(sc["1"], 266.155, abs_tol=1e-3)
    assert isclose(sc["2"], 4400.0, abs_tol=1e-3)
    assert isclose(sc["3"], 193.097, abs_tol=1e-3)
    assert isclose(ce.quantified_total(src)[0], 4859.252, abs_tol=1e-3)


def test_contractor_baselines_preserved_and_b_unavailable():
    inv = inventory_by_supplier()
    a = next(v for k, v in inv.items() if k.endswith("A"))
    c = next(v for k, v in inv.items() if k.endswith("C"))
    b = next(v for k, v in inv.items() if k.endswith("B"))
    assert isclose(a["emissions_t"], 121.956, abs_tol=1e-3)
    assert isclose(c["emissions_t"], 71.141, abs_tol=1e-3)
    assert b["emissions_t"] is None and b["tonne_km"] is None   # UNAVAILABLE, not zero


def test_record_completeness_48_of_60():
    audit, _ = ce.run()
    src = ce.by_source(audit)
    assert sum(s["records_with_data"] for s in src) == 48
    assert sum(s["records"] for s in src) == 60


def test_phase11_runs_do_not_change_inventory(d):
    before = inventory_by_supplier()
    _run(d, wtp=10_000, hypothetical={next(k for k in before if k.endswith("B")): 200.0})
    apply_offer_edits(d, offers_table(d))
    assert inventory_by_supplier() == before


def test_register_and_phase4_inputs_not_modified(d):
    reg = _sha(REGISTER)
    p4 = {n: _sha(ce.PHASE4 / "data" / n) for n in ce.INPUT_FILES}
    _run(d)
    apply_offer_edits(d, offers_table(d) + [{"offer_id": "X", "supplier_id": "NEW", "lane_id": "L1"}])
    assert _sha(REGISTER) == reg
    assert {n: _sha(ce.PHASE4 / "data" / n) for n in ce.INPUT_FILES} == p4
    man = json.loads((HERE / "outputs" / "run_manifest.json").read_text())
    for n, h in man["inputs"].items():
        assert _sha(HERE / n) == h if (HERE / n).exists() else True


def test_hypothetical_b_is_labelled_sensitivity_and_never_inserted(res, data):
    inv = inventory_by_supplier()
    bid = next(k for k in inv if k.endswith("B"))
    r = _run(copy.deepcopy(data), hypothetical={bid: 200.0})
    b = r["engagement"]["SUP-CARR-B"]
    assert b.impact.basis.startswith("SENSITIVITY") and b.impact.score is not None
    assert any("hypothetical" in x for x in b.limitations)
    rows, summ = r["ledger"]["inventory"]
    assert "OF-B-L4" in summ["unavailable_rows"]
    assert isclose(summ["baseline_t"], 193.097, abs_tol=1e-3)     # B not added to any ledger total


# ================================================================== 2 carbon
def test_tonne_km_and_emissions_basic():
    assert sce.tonne_km(2500, 400) == 1_000_000
    assert isclose(sce.emissions_t(1_000_000, 0.12), 120.0)
    assert isclose(sce.emissions_t(1_000_000, 0.08), 80.0)


def test_demo_carbon_comparison(data):
    c = sce.compare(data, "OF-A-L1", "OF-D-L1")
    assert c.comparable
    assert isclose(c.current_emissions_t, 120.0) and isclose(c.alternative_emissions_t, 80.0)
    assert isclose(c.avoided_t, 40.0) and isclose(c.reduction_pct, 1 / 3)
    assert c.direction == "reduction"
    assert isclose(c.current_intensity, 0.12) and isclose(c.alternative_intensity, 0.08)


def test_negative_activity_rejected_and_zero_activity_no_division():
    with pytest.raises(ValueError):
        sce.tonne_km(-1, 100)
    with pytest.raises(ValueError):
        sce.emissions_t(-5, 0.1)
    assert sce.intensity(10.0, 0) is None


def test_ef_unit_normalisation_and_invalid_unit():
    assert isclose(normalise_ef(80, "gCO2e/tonne-km"), 0.08)
    assert isclose(normalise_ef(0.12, "kgCO2e/tonne-km"), 0.12)
    with pytest.raises(UnitError):
        normalise_ef(0.1, "kgCO2e/litre")


def test_parse_number_edge_cases():
    issues = []
    assert parse_number(None, "x", issues) is None and not issues
    assert parse_number("abc", "x", issues) is None and issues[-1].level == "ERROR"
    assert parse_number(-3, "x", issues) is None and "negative" in issues[-1].message
    assert parse_number("12.5", "x", issues) == 12.5


def test_boundary_mismatch_not_comparable(d):
    d.offers["OF-D-L1"].ef.boundary = "Well-to-wheel"
    c = sce.compare(d, "OF-A-L1", "OF-D-L1")
    assert not c.comparable and any("boundary" in x for x in c.limitations)
    assert c.avoided_t is None


def test_different_lanes_are_not_substitutes(data):
    c = sce.compare(data, "OF-C-L2", "OF-D-L1")
    assert not c.comparable and any("different lanes" in x for x in c.limitations)


def test_partial_substitution(data):
    c = sce.compare(data, "OF-A-L1", "OF-D-L1", share=0.4)
    assert isclose(c.substituted_tonne_km, 400_000) and isclose(c.avoided_t, 16.0)
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", c.avoided_t, share=0.4)
    assert isclose(a.annual_cost_difference, 8_000)
    assert isclose(a.npv, -91_773.88, abs_tol=0.01)          # register Validation Cases, 40% case


def test_higher_emission_alternative(d):
    d.offers["OF-D-L1"].ef.value_kg_per_tkm = 0.13
    c = sce.compare(d, "OF-A-L1", "OF-D-L1")
    assert c.direction == "increase" and c.avoided_t < 0
    a = sfe.appraise(d, "OF-A-L1", "OF-D-L1", c.avoided_t)
    assert a.annual_cost_per_t is None and "HIGHER" in a.note


def test_decision_vs_inventory_basis_is_methodology_difference(res):
    dec, inv = res["ledger"]["decision"][1], res["ledger"]["inventory"][1]
    assert isclose(inv["baseline_t"], 193.097, abs_tol=1e-3)
    l1 = next(r for r in res["ledger"]["decision"][0] if r.offer_id == "OF-A-L1")
    l1i = next(r for r in res["ledger"]["inventory"][0] if r.offer_id == "OF-A-L1")
    assert isclose(l1.baseline_t - l1i.baseline_t, 1_000_000 * (0.12 - 0.10163) / 1000, abs_tol=0.01)  # 18.37 t
    assert dec["baseline_t"] > inv["baseline_t"]            # reported separately, inventory untouched


# ================================================================== 3 screening & data quality
def test_screening_classes(res):
    el = {k: v.eligibility for k, v in res["screening"].items()}
    assert el == {"OF-A-L1": "Conditionally eligible", "OF-A-L3": "Conditionally eligible",
                  "OF-C-L2": "Conditionally eligible", "OF-B-L4": "Insufficient evidence",
                  "OF-D-L1": "Conditionally eligible", "OF-E-L1": "Ineligible"}


def test_fourteen_checks_with_auto_derivation(res):
    for r in res["screening"].values():
        assert [c.check_no for c in r.checks] == list(range(1, 15))
        assert all(c.derived for c in r.checks if 3 <= c.check_no <= 10)


def test_safety_gate_failure_is_ineligible(res):
    e = res["screening"]["OF-E-L1"]
    assert e.eligibility == "Ineligible" and any("check 2" in x for x in e.blocking_reasons)


def test_capacity_shortfall_is_ineligible(d):
    d.offers["OF-D-L1"].capacity_t = 2000
    r = sse.screen_offer(d, "OF-D-L1")
    assert r.eligibility == "Ineligible"
    assert next(c for c in r.checks if c.check_no == 3).status == "Fail"


def test_incompatible_period_blocks_comparison(d):
    d.offers["OF-D-L1"].period = "2024"
    r = sse.screen_offer(d, "OF-D-L1")
    assert next(c for c in r.checks if c.check_no == 6).status == "Fail"
    assert r.eligibility in ("Insufficient evidence", "Ineligible")


def test_boundary_failure_conditional_not_carbon_comparable(d):
    d.offers["OF-D-L1"].ef.boundary = "Well-to-wheel"
    r = sse.screen_offer(d, "OF-D-L1")
    assert r.eligibility == "Conditionally eligible" and not r.carbon_comparable


def test_verification_policy_switch(d):
    assert sse.screen_offer(d, "OF-D-L1").eligibility == "Conditionally eligible"
    d.settings.verification_mandatory = True
    r = sse.screen_offer(d, "OF-D-L1")
    assert r.eligibility == "Insufficient evidence" and any("verification" in x for x in r.missing_material)


def test_data_quality_dimensions(data):
    a, b = sse.data_quality(data, "SUP-CARR-A"), sse.data_quality(data, "SUP-CARR-B")
    assert len(a.dimensions) == 7 and a.usable_estimate and "usable" in a.note
    assert not b.usable_estimate and "UNAVAILABLE" in b.note and b.score < 10   # rated mostly Missing (= 0)


def test_equivalence_reports_limitations(d):
    d.offers["OF-D-L1"].mass_t = 1000
    e = sse.equivalence(d, "OF-A-L1", "OF-D-L1", 1.0)
    assert not e.equivalent and any("less than" in x for x in e.limitations)


# ================================================================== 4 engagement matrix
def test_contractor_a_priority_1(res):
    a = res["engagement"]["SUP-CARR-A"]
    assert isclose(a.impact.score, 70.53, abs_tol=0.01) and isclose(a.maturity.score, 39.29, abs_tol=0.01)
    assert a.quadrant == "Q1" and a.priority == 1


def test_contractor_c_maintain(res):
    c = res["engagement"]["SUP-CARR-C"]
    assert c.quadrant == "Q4" and c.priority == 4


def test_contractor_b_unclassified(res):
    b = res["engagement"]["SUP-CARR-B"]
    assert b.quadrant == see.UNCLASSIFIED and b.label == see.UNCLASSIFIED_LABEL
    assert b.impact.score is None and b.plan          # default data-request actions exist


@pytest.mark.parametrize("impact,maturity,q", [(80, 20, "Q1"), (80, 80, "Q2"), (20, 20, "Q3"), (20, 80, "Q4"),
                                               (50, 50, "Q2"), (49.99, 50, "Q4"), (None, 60, "UNCLASSIFIED"),
                                               (60, None, "UNCLASSIFIED")])
def test_quadrant_classification(data, impact, maturity, q):
    assert see.classify_engagement_quadrant(impact, maturity, data.matrix) == q


def test_editable_thresholds_change_classification(d):
    d.matrix.impact_threshold = 80
    r = _run(d)
    assert r["engagement"]["SUP-CARR-A"].quadrant == "Q3"
    d.matrix.impact_threshold, d.matrix.maturity_threshold = 30, 30
    r = _run(d)
    assert r["engagement"]["SUP-CARR-A"].quadrant == "Q2" and r["engagement"]["SUP-CARR-C"].quadrant == "Q2"


def test_minimum_maturity_indicators(d):
    d.matrix.min_maturity_indicators = 8
    m = see.assess_supplier_maturity(d, "SUP-CARR-A")
    assert m.score is None and m.status == see.UNCLASSIFIED


def test_prospective_suppliers_not_in_matrix(res, data):
    assert set(res["engagement"]) == {s for s, x in data.suppliers.items() if x.is_existing}
    imp = see.calculate_supplier_impact(data, inventory_by_supplier(), "ALT-D")
    assert imp.status == "NOT APPLICABLE"


def test_matrix_never_triggers_replacement(d):
    for ti, tm in ((0, 100), (100, 0), (50, 50), (90, 90)):
        d.matrix.impact_threshold, d.matrix.maturity_threshold = ti, tm
        r = _run(d)
        assert all(x.action != sre.SWITCH for x in r["recommendations"].values())
        for e in r["engagement"].values():
            assert not any("replace" in a.lower() or "terminate" in a.lower() for a in e.actions)


def test_action_plan_fields_and_deadlines(res):
    plan = [a for e in res["engagement"].values() for a in e.plan]
    need = {"action_id", "responsible_role", "proposed_deadline", "required_evidence", "success_indicator", "status",
            "follow_up_date", "verification_method", "deadline_basis"}
    assert plan and all(need <= set(a) for a in plan)
    assert {a["deadline_days"] for a in plan} <= {30, 60, 90}
    assert all("management assumption" in a["deadline_basis"] for a in plan)
    assert len(res["engagement"]["SUP-CARR-B"].plan) >= 5


# ================================================================== 5 financial
def test_demo_npv_matches_workbook_validation_case(data):
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 40.0)
    assert isclose(a.npv, data.validation_cases["Switching NPV (Excel PV function)"], abs_tol=0.01)
    assert isclose(a.npv, -184_434.70, abs_tol=0.01)
    assert abs(a.npv - (-184_443)) < 10      # brief's rounded figure differs by SAR 8.30 (reported)


def test_demo_cost_metrics(data):
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 40.0)
    assert isclose(a.annual_cost_difference, 20_000) and isclose(a.annual_cost_per_t, 500)
    assert isclose(a.breakeven_alternative_cost, 296_114.86, abs_tol=0.01)
    assert isclose(a.lifetime_discounted_cost_per_t, 597.128, abs_tol=0.001)
    assert a.payback_years is None and a.breakeven_switching_cost is None


def test_discount_rate_change(data):
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 40.0, rate=0.0)
    assert isclose(a.npv, -30_000 - 20_000 * 10)


def test_switching_cost_change(data):
    a0 = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 40.0)
    a1 = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 40.0, switching_cost=60_000)
    assert isclose(a1.npv, a0.npv - 30_000)


def test_horizon_change(data):
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 40.0, horizon=5)
    af = sum(1 / 1.05 ** t for t in range(1, 6))
    assert isclose(a.npv, -30_000 - 20_000 * af, abs_tol=0.01)


def test_cheaper_alternative_positive_npv_and_payback(data):
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 40.0, alt_cost=290_000)
    assert a.npv > 0 and isclose(a.payback_years, 3.0) and isclose(a.breakeven_switching_cost, 10_000 * a.annuity_factor)


def test_missing_cost_gives_incomplete_appraisal(d):
    d.offers["OF-D-L1"].annual_cost = None
    a = sfe.appraise(d, "OF-A-L1", "OF-D-L1", 40.0)
    assert not a.complete and a.npv is None and "missing" in a.note


def test_zero_avoided_cost_per_tonne_undefined(data):
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 0.0)
    assert a.annual_cost_per_t is None and a.lifetime_discounted_cost_per_t is None


def test_willingness_to_pay_check(data):
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", 40.0)
    assert sfe.within_willingness_to_pay(a, None) is None
    assert sfe.within_willingness_to_pay(a, 600) is True and sfe.within_willingness_to_pay(a, 500) is False


# ================================================================== 6 scoring
def test_default_weights_normalised(data):
    w = ssc.validate_weights(data.scoring.weights)
    assert isclose(sum(w.values()), 1.0) and w == pytest.approx(ssc.DEFAULT_WEIGHTS)


@pytest.mark.parametrize("bad", [{"carbon_performance": -1}, {"carbon_performance": "abc"}, {"carbon_performance": True},
                                 "zero", "missing", {"unknown": 1}])
def test_invalid_weights_rejected(bad):
    w = dict(ssc.DEFAULT_WEIGHTS)
    if bad == "zero":
        w = {k: 0 for k in w}
    elif bad == "missing":
        w.pop("carbon_performance")
    else:
        w.update(bad)
    with pytest.raises(WeightError):
        ssc.validate_weights(w)


def test_unnormalised_weights_equivalent(res, data):
    s1 = ssc.score_lane(data, "L1", res["screening"])
    s2 = ssc.score_lane(data, "L1", res["screening"], {"carbon_performance": 70, "total_lifecycle_cost": 50,
                                                       "operational_reliability": 40, "data_quality_transparency": 30,
                                                       "improvement_capability": 10})
    assert [r.total for r in s1.rows] == pytest.approx([r.total for r in s2.rows])


def test_lane_l1_scores_and_ranking(res):
    s = res["scoring"]["L1"]
    t = {r.offer_id: r for r in s.rows}
    assert s.ranking_status == "DEFINITIVE"
    assert isclose(t["OF-D-L1"].total, 56.20, abs_tol=0.01) and t["OF-D-L1"].rank == 1
    assert isclose(t["OF-A-L1"].total, 44.96, abs_tol=0.01) and t["OF-A-L1"].rank == 2
    for r in s.rows:
        if r.total is not None:
            assert isclose(sum(r.contributions.values()), r.total)


def test_eligibility_overrides_score(res):
    t = {r.offer_id: r for r in res["scoring"]["L1"].rows}
    e = t["OF-E-L1"]
    assert e.total > t["OF-D-L1"].total and e.rank is None and "NOT ranked" in e.status


def test_missing_mandatory_input_suppresses_ranking(d):
    d.offers["OF-D-L1"].on_time_pct = None
    s = ssc.score_lane(d, "L1", sse.screen_all(d))
    t = {r.offer_id: r for r in s.rows}
    assert s.ranking_status == "SUPPRESSED" and t["OF-D-L1"].total is None
    assert all(r.rank is None for r in s.rows)


def test_weight_change_reverses_ranking(res, data):
    w = {"carbon_performance": 0.10, "total_lifecycle_cost": 0.60, "operational_reliability": 0.10,
         "data_quality_transparency": 0.15, "improvement_capability": 0.05}
    assert ssc.score_lane(data, "L1", res["screening"]).top() == "OF-D-L1"
    assert ssc.score_lane(data, "L1", res["screening"], w).top() == "OF-A-L1"
    rows, switch = ssc.weight_sensitivity(data, "L1", res["screening"])
    assert len(rows) == 10 and switch is not None and switch["top_from"] == "OF-D-L1"


# ================================================================== 7 recommendations
def test_default_recommendations(res):
    r = {k: (v.action, v.rule_id) for k, v in res["recommendations"].items()}
    assert r == {"L1": (sre.NONE, "R6"), "L2": (sre.RETAIN, "R3"), "L3": (sre.IMPROVE, "R3"), "L4": (sre.DATA, "R2")}


def test_consider_switching_within_willingness_to_pay(d):
    r = _run(d, wtp=600)["recommendations"]["L1"]
    assert r.action == sre.SWITCH and r.rule_id == "R7" and r.human_approval_required
    assert any("procurement approval" in s for s in r.next_steps)


def test_cost_above_willingness_to_pay(d):
    r = _run(d, wtp=500)["recommendations"]["L1"]
    assert r.action == sre.IMPROVE and r.rule_id == "R8"


def test_consider_switching_when_npv_non_negative(d):
    d.offers["OF-D-L1"].annual_cost = 290_000
    r = _run(d)["recommendations"]["L1"]
    assert r.action == sre.SWITCH and r.rule_id == "R5"


def test_not_robust_carbon_gives_no_defensible(d):
    d.offers["OF-D-L1"].annual_cost = 290_000
    d.uncertainty["alternative_ef"].high = 0.13          # unfavourable case: alternative dirtier
    r = _run(d)["recommendations"]["L1"]
    assert r.action == sre.NONE and r.rule_id == "R5u"


def test_ineligible_incumbent_escalates(d):
    for e in d.screening["OF-C-L2"]:
        if e.check_no == 2:
            e.status = "Fail"
    r = _run(d)["recommendations"]["L2"]
    assert r.action == sre.INELIGIBLE and r.rule_id == "R1" and r.human_approval_required


def test_financial_only_gain_needs_judgement(d):
    d.offers["OF-D-L1"].ef.value_kg_per_tkm = 0.13
    d.offers["OF-D-L1"].annual_cost = 250_000
    r = _run(d)["recommendations"]["L1"]
    assert r.action == sre.NONE and r.rule_id == "R9"


def test_ineligible_alternative_excluded_from_recommendation(res):
    r = res["recommendations"]["L1"]
    assert r.best_alternative == "OF-D-L1" and any(x.startswith("OF-E-L1: INELIGIBLE") for x in r.excluded_alternatives)


def test_recommendation_explanations_complete(res):
    for r in res["recommendations"].values():
        assert r.rule and r.material_assumptions and r.next_steps and r.approval_note and r.synthetic
        assert r.carbon_implications and r.financial_implications
        assert r.action in sre.APPROVAL and r.human_approval_required == sre.APPROVAL[r.action]
    l1 = res["recommendations"]["L1"]
    assert l1.supporting_evidence and l1.sensitivity and l1.operational_implications and l1.data_limitations
    assert "NOT SET" in " ".join(l1.material_assumptions)


def test_six_action_types_reachable():
    assert {sre.RETAIN, sre.IMPROVE, sre.SWITCH, sre.DATA, sre.INELIGIBLE, sre.NONE} == set(sre.APPROVAL)


def test_recommendation_robustness(res):
    rows, robust = res["uncertainty"]["robustness"]
    assert not robust and rows[0]["case"] == "Base"
    assert any(x["action"] == sre.SWITCH for x in rows)       # alternative cost 290k case


def test_missing_ef_on_incumbent_does_not_crash(d):
    d.offers["OF-A-L3"].ef = None
    d.offers["OF-D-L1"].capacity_t = None
    r = _run(d)["recommendations"]
    assert r["L3"].action in (sre.DATA, sre.IMPROVE, sre.RETAIN)


# ================================================================== 8 uncertainty & integration
def test_joint_cases(res):
    j = res["uncertainty"]["joint"]
    assert isclose(j["Base"]["avoided_t"], 40) and isclose(j["Base"]["npv"], -184_434.70, abs_tol=0.01)
    assert isclose(j["Unfavourable to switching"]["avoided_t"], 10, abs_tol=1e-6)
    assert isclose(j["Favourable to switching"]["avoided_t"], 70, abs_tol=1e-6)
    assert j["Unfavourable to switching"]["npv"] < j["Base"]["npv"] < j["Favourable to switching"]["npv"]


def test_one_way_base_equals_deterministic(res):
    u = res["uncertainty"]
    assert isclose(u["base"]["npv"], -184_434.70, abs_tol=0.01)
    vs = {r["variable"] for r in u["one_way"]}
    assert {"alternative_ef", "alternative_annual_cost", "switching_cost", "current_ef"} <= vs


def test_breakevens(res):
    b = res["uncertainty"]["breakevens"]
    assert isclose(b["financial_breakeven_alternative_annual_cost_sar"], 296_114.86, abs_tol=0.01)
    assert b["carbon_breakeven_alternative_ef_kg_per_tkm"] == 0.12
    assert isclose(b["willingness_to_pay_needed_sar_per_t"], 597.128, abs_tol=0.001)


def test_two_variable_grid_zones(res):
    g = res["uncertainty"]["grid"]
    zones = {r["zone"] for r in g}
    assert len(g) == 9 * 8
    assert {"No carbon benefit", "Carbon and financial gain", "Carbon gain at net cost (management judgement)"} <= zones


def test_no_ranking_reversal_in_parameter_ranges(res):
    assert all(not x["reversal"] for x in res["uncertainty"]["ranking_reversals"])


def test_ledger_inventory_basis_matches_workstream_c(res):
    s = res["ledger"]["inventory"][1]
    wc = next(x for x in se.run_all()["comparison"] if x["scenario"] == "Known freight activity optimization")
    # Workstream C sums record-level (rounded) emissions; the ledger multiplies annual tonne-km: <= 1e-5 t apart
    assert isclose(s["efficiency_avoided_t"], wc["avoided_t"], abs_tol=1e-5)
    assert isclose(s["efficiency_avoided_t"], 28.965, abs_tol=1e-3)


def test_ledger_no_double_counting(res):
    s = res["ledger"]["decision"][1]
    assert isclose(s["baseline_t"], 211.467, abs_tol=1e-3)
    assert isclose(s["efficiency_avoided_t"], 0.15 * s["baseline_t"]) and isclose(s["switching_avoided_t"], 40.0)
    assert isclose(s["overlap_t"], 0.15 * 40.0, abs_tol=1e-9)        # efficiency share of the switched activity
    assert isclose(s["combined_avoided_t"], s["efficiency_avoided_t"] + 40.0 - 6.0, abs_tol=1e-9)
    assert s["combined_avoided_t"] < s["sum_of_separate_avoided_t"]
    row = next(r for r in res["ledger"]["decision"][0] if r.offer_id == "OF-A-L1")
    assert isclose(row.combined_t, 1_000_000 * 0.85 * 0.08 / 1000)   # adjusted activity x alternative EF


def test_validation_cases_cached_values_match_engine(data):
    vc = data.validation_cases
    c = sce.compare(data, "OF-A-L1", "OF-D-L1")
    a = sfe.appraise(data, "OF-A-L1", "OF-D-L1", c.avoided_t)
    for key, val in (("Avoided emissions", c.avoided_t), ("Percentage reduction", c.reduction_pct),
                     ("Annual cost difference (alternative - current)", a.annual_cost_difference),
                     ("Incremental annual cost per tonne avoided", a.annual_cost_per_t),
                     ("Annuity factor", a.annuity_factor), ("Switching NPV (closed form)", a.npv),
                     ("Break-even alternative annual cost", a.breakeven_alternative_cost),
                     ("Discounted lifetime cost per tonne", a.lifetime_discounted_cost_per_t)):
        assert isclose(vc[key], val, rel_tol=1e-9, abs_tol=1e-6), key


def test_register_loads_clean(data):
    assert not [i for i in data.issues if i.level == "ERROR"]
    assert len(data.input_version) == 64
    wb = openpyxl.load_workbook(REGISTER)
    assert len(wb.sheetnames) == 15 and wb.sheetnames[0] == "README"


def test_register_duplicate_and_unsupported_mode_reported(tmp_path):
    p = _mutate(tmp_path, "Supplier Profiles", "ALT-E", "transport_mode", "hovercraft")
    d = load_register(p)
    assert any("unsupported transport mode" in i.message for i in d.issues)
    p = _mutate(tmp_path, "Freight Activity", "OF-E-L1", "offer_id", "OF-D-L1")
    d = load_register(p)
    assert any("duplicate offer_id" in i.message for i in d.issues)


def test_register_non_numeric_activity_reported(tmp_path):
    p = _mutate(tmp_path, "Freight Activity", "OF-D-L1", "freight_mass_t", "lots")
    d = load_register(p)
    assert any(i.level == "ERROR" and "non-numeric" in i.message for i in d.issues)
    assert d.offers["OF-D-L1"].tonne_km is None
    r = _run(d)
    assert r["recommendations"]["L1"].best_alternative is None      # incomplete alternative excluded


def test_dashboard_edits_user_added_supplier(data):
    rows = offers_table(data) + [{"offer_id": "OF-F-L1", "supplier_id": "ALT-F", "supplier_name": "Carrier F",
                                  "lane_id": "L1", "freight_mass_t": 2500, "avg_distance_km": 400,
                                  "reporting_period": "2025", "offered_capacity_t_per_year": 3000,
                                  "ef_kgco2e_per_tkm": 0.05, "ef_boundary": "Tank-to-wheel",
                                  "annual_service_cost_sar": 330000, "on_time_delivery_pct": 96,
                                  "safety_legal_status": "Pass"}]
    d = apply_offer_edits(data, rows)
    assert d.suppliers["ALT-F"].relationship == "User-added"
    r = sse.screen_offer(d, "OF-F-L1")
    assert next(c for c in r.checks if c.check_no == 11).status == "Missing"     # never silently passed
    assert "OF-F-L1" not in data.offers                                           # original untouched


def test_dashboard_edit_errors(data):
    rows = offers_table(data)
    rows.append(dict(rows[0]))                        # duplicate ID
    rows.append({"offer_id": "OF-X", "supplier_id": "ALT-X", "lane_id": "L9"})   # unknown lane
    rows[4] = {**rows[4], "annual_service_cost_sar": -10}
    d = apply_offer_edits(data, rows)
    msgs = " ".join(i.message for i in d.issues)
    assert "duplicate offer_id" in msgs and "unknown lane_id" in msgs and "negative" in msgs
    assert d.offers["OF-D-L1"].annual_cost is None


def test_deleted_offer_removed_in_session_copy(data):
    d = apply_offer_edits(data, [r for r in offers_table(data) if r["offer_id"] != "OF-E-L1"])
    assert "OF-E-L1" not in d.offers and "OF-E-L1" in data.offers


def test_audit_records_complete_and_no_invented_verification(res, data):
    au = sa.audit_records(res)
    need = {"input_version", "engine_version", "supplier_id", "source", "reporting_period", "calculation_method",
            "emission_factor", "model_assumptions", "recommendation_rule", "evidence_status", "ef_verification"}
    assert au and all(need <= set(r) for r in au)
    allowed = {o.ef.verification for o in data.offers.values() if o.ef} | {None}
    assert {r["ef_verification"] for r in au} <= allowed
    assert not any("Verified - third party" == r["ef_verification"] for r in au)


def test_monitoring_scorecard_projected_vs_verified(res):
    sc = sa.monitoring_scorecard(res)
    assert {r["supplier_id"] for r in sc} == set(res["engagement"])
    assert all(r["verified_actual_reduction_tco2e"].startswith("none") for r in sc)


def test_phase11_outputs_reproducible(tmp_path):
    sa.main(tmp_path / "a")
    sa.main(tmp_path / "b")
    a = json.loads((tmp_path / "a" / "run_manifest_p11.json").read_text())
    b = json.loads((tmp_path / "b" / "run_manifest_p11.json").read_text())
    assert a["outputs"] == b["outputs"] and len(a["outputs"]) >= 20


# ================================================================== 9 Workstream C unchanged
def test_workstream_c_outputs_unchanged(tmp_path):
    committed = json.loads((HERE / "outputs" / "workstream_c" / "run_manifest_c.json").read_text())
    assert rwc.main(tmp_path) == 0
    fresh = json.loads((tmp_path / "run_manifest_c.json").read_text())
    assert fresh["outputs"] == committed["outputs"]


# ================================================================== 10 Streamlit smoke
@pytest.fixture(scope="module")
def app():
    from streamlit.testing.v1 import AppTest
    return AppTest.from_file(str(HERE / "app.py"), default_timeout=180).run()


def _actions(at):
    import re
    out = {}
    for m in at.markdown:
        g = re.search(r">([A-Z /]+)</span><span style='font-weight:600'>(L\d)", m.value)
        if g:
            out[g.group(2)] = g.group(1)
    return out


def test_app_nine_tabs_no_exceptions(app):
    assert not app.exception
    labels = [t.label for t in app.tabs][:9]
    assert labels == ["Hotspots", "Monthly trends", "Data confidence", "Emission factors", "Audit trail",
                      "Reduction scenarios", "Financial evaluation", "Decision intelligence",
                      "Supplier Decision & Engagement"]


def test_app_objective_and_sections(app):
    import phase11_tab
    assert any(phase11_tab.OBJECTIVE in m.value for m in app.markdown)
    subs = [s.value for s in app.subheader]
    for s in ("Supplier overview", "Screening & data quality", "Supplier engagement priority matrix",
              "Alternatives comparison", "Carbon & financial trade-off", "Explainable recommendations",
              "Engagement action plans", "Monitoring, verification & audit"):
        assert s in subs


def test_app_recommendations_and_interactions():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(HERE / "app.py"), default_timeout=180).run()
    assert _actions(at) == {"L1": sre.NONE, "L2": sre.RETAIN, "L3": sre.IMPROVE, "L4": sre.DATA}
    at.checkbox(key="p11_wtp_set").check().run()
    at.number_input(key="p11_wtp").set_value(600).run()
    assert _actions(at)["L1"] == sre.SWITCH
    at.number_input(key="p11_w_carbon_performance").set_value(-5).run()
    assert any("Invalid weights" in e.value for e in at.error)
    at.slider(key="p11_imp_thr").set_value(80).run()
    q = next(d.value for d in at.dataframe if {"Quadrant", "Impact"} <= set(d.value.columns))
    assert dict(zip(q["Supplier"], q["Quadrant"]))["SUP-CARR-A"] == "Q3"
    assert not at.exception
